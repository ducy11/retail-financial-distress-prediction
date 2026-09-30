"""Xuất Word (.docx) và Slide (.pptx) từ báo cáo markdown, không cần thêm phụ thuộc nặng.

- Word: dùng `python-docx` (đã có trong môi trường) để dựng tiêu đề, bảng, hình, code block.
- Slide: tự ghi OOXML tối giản bằng `zipfile` (máy không có `python-pptx` vẫn xuất được .pptx).

Lệnh: python -m scripts.export_office
      → docs/BAO-CAO.docx                  (BÁO CÁO HOÀN CHỈNH: 11 mục + 3 phụ lục, kèm bảng & hình)
        docs/BAO-CAO-slide.pptx            (slide tự động 14 mục, khớp artifact)
        docs/BAO-CAO-slide-bao-ve.pptx     (deck bảo vệ 11 slide)
        docs/BAO-CAO-slide-bao-ve.docx     (cùng nội dung deck bảo vệ, dạng Word để DỰNG SLIDE)
        docs/bo-tai-lieu-bao-ve.docx       (bộ tài liệu bảo vệ: factsheet + dàn 11 slide + 8 Q&A)
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Tuple

from forecasting.config import DOCS_DIR, ensure_utf8_stdio

#: Các pattern markdown cần xử lý.
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_CODE = re.compile(r"`([^`]+)`")
_IMG = re.compile(r"!\[(.*?)\]\((.+?)\)")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_TABLE_ROW = re.compile(r"^\|(.+)\|$")
_SEP_ROW = re.compile(r"^\|[\s:\-|]+\|$")
_BULLET = re.compile(r"^[-*]\s+(.*)$")
_NUMBERED = re.compile(r"^\d+\.\s+(.*)$")


def _clean(text: str) -> str:
    """Bỏ ký hiệu markdown khi không cần định dạng (ví dụ trong ô bảng)."""
    text = _BOLD.sub(r"\1", text)
    text = _CODE.sub(r"\1", text)
    return text.replace("\\|", "|").strip()


def _add_runs(par, text: str) -> None:
    """Thêm run vào paragraph, tôn trọng `**bold**` và `` `code` ``."""
    for token in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            par.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = par.add_run(token[1:-1])
            run.font.name = "Consolas"
        else:
            par.add_run(token)


def _parse_table(lines: List[str]) -> List[List[str]]:
    """Đọc khối bảng markdown thành list các hàng (bỏ dòng phân cách)."""
    rows: List[List[str]] = []
    for line in lines:
        if _SEP_ROW.match(line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append([_clean(c) for c in cells])
    return rows


def _md_blocks(md: str) -> List[Tuple[str, Any]]:
    """Tách markdown thành block: heading / paragraph / bullet / numbered / table / image / code."""
    blocks: List[Tuple[str, Any]] = []
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            i += 1
            code: List[str] = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(code)))
        elif _HEADING.match(line):
            m = _HEADING.match(line)
            blocks.append(("heading", (len(m.group(1)), _clean(m.group(2)))))
        elif _IMG.search(line):
            m = _IMG.search(line)
            blocks.append(("image", (m.group(1), m.group(2))))
        elif _TABLE_ROW.match(line):
            table_lines = []
            while i < len(lines) and _TABLE_ROW.match(lines[i]):
                table_lines.append(lines[i])
                i += 1
            blocks.append(("table", _parse_table(table_lines)))
            continue
        elif _BULLET.match(line):
            blocks.append(("bullet", _clean(_BULLET.match(line).group(1))))
        elif _NUMBERED.match(line):
            blocks.append(("numbered", _clean(_NUMBERED.match(line).group(1))))
        elif line.strip():
            blocks.append(("paragraph", line.strip()))
        i += 1
    return blocks


def build_docx(md_path: Path, out_path: Path) -> Dict[str, Any]:
    """Dựng file .docx từ báo cáo markdown (kèm bảng và hình)."""
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    counts = {"heading": 0, "paragraph": 0, "table": 0, "image": 0, "missing_image": 0,
              "bullet": 0, "code": 0}
    for kind, payload in _md_blocks(md_path.read_text(encoding="utf-8")):
        if kind == "heading":
            level, text = payload
            doc.add_heading(text, level=min(level, 4))
            counts["heading"] += 1
        elif kind == "paragraph":
            _add_runs(doc.add_paragraph(), payload)
            counts["paragraph"] += 1
        elif kind == "bullet":
            _add_runs(doc.add_paragraph(style="List Bullet"), payload)
            counts["bullet"] += 1
        elif kind == "numbered":
            _add_runs(doc.add_paragraph(style="List Number"), payload)
            counts["bullet"] += 1
        elif kind == "code":
            run = doc.add_paragraph().add_run(payload)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            counts["code"] += 1
        elif kind == "table":
            rows = payload
            if not rows:
                continue
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                cells = table.add_row().cells
                for c, value in enumerate(row[:len(rows[0])]):
                    par = cells[c].paragraphs[0]
                    _add_runs(par, value)
                    if r == 0:
                        for run in par.runs:
                            run.bold = True
            counts["table"] += 1
        elif kind == "image":
            alt, rel = payload
            image_path = (md_path.parent / rel).resolve()
            if image_path.exists():
                doc.add_picture(str(image_path), width=Inches(6.0))
                doc.paragraphs[-1].alignment = 1  # căn giữa
                counts["image"] += 1
            else:
                doc.add_paragraph(f"[Thiếu hình: {rel} ({alt})]")
                counts["missing_image"] += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return counts


# ---------------------------------------------------------------------------
# PPTX: tự ghi OOXML tối giản (không cần python-pptx)
# ---------------------------------------------------------------------------
_NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
       'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')

EMU_PER_INCH = 914400
SLIDE_W = 12192000   # 13,333 inch
SLIDE_H = 6858000    # 7,5 inch

_SPTREE_EMPTY = (
    '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
    '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
)
_NV_GRP = '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'

CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Default Extension="png" ContentType="image/png"/>'
    '<Default Extension="jpg" ContentType="image/jpeg"/>'
    '<Default Extension="jpeg" ContentType="image/jpeg"/>'
    '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
    '<Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/>'
    '<Override PartName="/ppt/slideLayouts/slideLayout1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"/>'
    '<Override PartName="/ppt/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>'
    '{slides}'
    '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
    '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
    '</Types>'
)

ROOT_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>'
    '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
    '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>'
    '</Relationships>'
)

THEME = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Office">'
    '<a:themeElements>'
    '<a:clrScheme name="Office">'
    '<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>'
    '<a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>'
    '<a:dk2><a:srgbClr val="44546A"/></a:dk2><a:lt2><a:srgbClr val="E7E6E6"/></a:lt2>'
    '<a:accent1><a:srgbClr val="4472C4"/></a:accent1><a:accent2><a:srgbClr val="ED7D31"/></a:accent2>'
    '<a:accent3><a:srgbClr val="A5A5A5"/></a:accent3><a:accent4><a:srgbClr val="FFC000"/></a:accent4>'
    '<a:accent5><a:srgbClr val="5B9BD5"/></a:accent5><a:accent6><a:srgbClr val="70AD47"/></a:accent6>'
    '<a:hlink><a:srgbClr val="0563C1"/></a:hlink><a:folHlink><a:srgbClr val="954F72"/></a:folHlink>'
    '</a:clrScheme>'
    '<a:fontScheme name="Office"><a:majorFont><a:latin typeface="Calibri Light"/>'
    '<a:ea typeface=""/><a:cs typeface=""/></a:majorFont>'
    '<a:minorFont><a:latin typeface="Calibri"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont>'
    '</a:fontScheme>'
    '<a:fmtScheme name="Office">'
    '<a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
    '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
    '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:fillStyleLst>'
    '<a:lnStyleLst><a:ln w="9525"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln>'
    '<a:ln w="12700"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln>'
    '<a:ln w="19050"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln></a:lnStyleLst>'
    '<a:effectStyleLst><a:effectStyle><a:effectLst/></a:effectStyle>'
    '<a:effectStyle><a:effectLst/></a:effectStyle><a:effectStyle><a:effectLst/></a:effectStyle>'
    '</a:effectStyleLst>'
    '<a:bgFillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
    '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
    '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:bgFillStyleLst>'
    '</a:fmtScheme></a:themeElements><a:objectDefaults/><a:extraClrSchemeLst/></a:theme>'
)

SLIDE_MASTER = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<p:sldMaster {_NS}><p:cSld><p:spTree>{_NV_GRP}{_SPTREE_EMPTY}</p:spTree></p:cSld>'
    '<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" '
    'accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" '
    'folHlink="folHlink"/>'
    '<p:sldLayoutIdLst><p:sldLayoutId id="2147483649" r:id="rId1"/></p:sldLayoutIdLst>'
    '<p:txStyles><p:titleStyle/><p:bodyStyle/><p:otherStyle/></p:txStyles></p:sldMaster>'
)

SLIDE_MASTER_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/>'
    '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" Target="../theme/theme1.xml"/>'
    '</Relationships>'
)

SLIDE_LAYOUT = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<p:sldLayout {_NS} type="blank" preserve="1"><p:cSld name="Blank">'
    f'<p:spTree>{_NV_GRP}{_SPTREE_EMPTY}</p:spTree></p:cSld>'
    '<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>'
)

SLIDE_LAYOUT_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="../slideMasters/slideMaster1.xml"/>'
    '</Relationships>'
)

CORE_PROPS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    '<dc:title>Đồ án dự báo suy giảm tài chính doanh nghiệp bán lẻ</dc:title>'
    '<dc:creator>scripts.export_office</dc:creator></cp:coreProperties>'
)

APP_PROPS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
    '<Application>python-ooxml-writer</Application></Properties>'
)


def _esc(text: str) -> str:
    """Escape XML cho nội dung text."""
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;"))


def _png_size(path: Path) -> Tuple[int, int]:
    """Đọc kích thước ảnh PNG từ header IHDR (không cần thư viện ảnh)."""
    with open(path, "rb") as f:
        head = f.read(24)
    if len(head) < 24 or head[:8] != b"\x89PNG\r\n\x1a\n":
        return 1, 1
    width = int.from_bytes(head[16:20], "big")
    height = int.from_bytes(head[20:24], "big")
    return max(1, width), max(1, height)


def _text_shape(shape_id: int, name: str, x: int, y: int, cx: int, cy: int,
                paragraphs: List[str], size: int = 1800, bold: bool = False,
                bullet: bool = False) -> str:
    """Một shape text (title hoặc body) cho slide."""
    runs = []
    for text in paragraphs:
        ppr = f'<a:pPr marL="{228600 if bullet else 0}" indent="{-228600 if bullet else 0}">' \
              + ('<a:buChar char="•"/>' if bullet else "<a:buNone/>") + '</a:pPr>'
        rpr = f'<a:rPr lang="vi-VN" sz="{size}" b="{1 if bold else 0}" dirty="0"/>'
        runs.append(f'<a:p>{ppr}<a:r>{rpr}<a:t>{_esc(text)}</a:t></a:r></a:p>')
    if not runs:
        runs.append('<a:p/>')
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{shape_id}" name="{_esc(name)}"/>'
        '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="square" rtlCol="0"/><a:lstStyle/>{"".join(runs)}</p:txBody></p:sp>'
    )


def _pic_shape(shape_id: int, rel_id: str, name: str, x: int, y: int, cx: int, cy: int) -> str:
    """Shape ảnh (p:pic) cho slide."""
    return (
        f'<p:pic><p:nvPicPr><p:cNvPr id="{shape_id}" name="{_esc(name)}"/><p:cNvPicPr/>'
        '<p:nvPr/></p:nvPicPr>'
        f'<p:blipFill><a:blip r:embed="{rel_id}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
    )


def _parse_slides(md_path: Path) -> List[Dict[str, Any]]:
    """Tách `docs/slide.md` thành list slide {title, bullets, image}."""
    slides: List[Dict[str, Any]] = []
    current: Dict[str, Any] | None = None
    for line in md_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            current = {"title": re.sub(r"^\d+\.\s*", "", line[3:].strip()),
                       "bullets": [], "image": None}
            slides.append(current)
        elif current is not None and _BULLET.match(line):
            current["bullets"].append(_clean(_BULLET.match(line).group(1)))
        elif current is not None and _IMG.search(line):
            current["image"] = _IMG.search(line).group(2)
    return slides


def _rels_xml(entries: List[Tuple[str, str, str]]) -> str:
    """Sinh XML relationships từ danh sách (id, type, target)."""
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + "".join(f'<Relationship Id="{i}" Type="{t}" Target="{target}"/>'
                      for i, t, target in entries)
            + '</Relationships>')


def build_pptx(md_path: Path, out_path: Path) -> int:
    """Dựng .pptx 16:9 từ `slide.md` (tiêu đề + bullet + hình), ghi OOXML bằng zipfile."""
    slides = _parse_slides(md_path)
    if not slides:
        raise ValueError(f"Không tìm thấy slide nào trong {md_path}")
    n = len(slides)

    content_types = CONTENT_TYPES.format(slides="".join(
        f'<Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        for i in range(1, n + 1)))
    presentation = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p:presentation {_NS} saveSubsetFonts="1">'
        '<p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>'
        + "".join(f'<p:sldId id="{255 + i}" r:id="rId{1 + i}"/>' for i in range(1, n + 1))
        + f'</p:sldIdLst><p:sldSz cx="{SLIDE_W}" cy="{SLIDE_H}"/>'
          f'<p:notesSz cx="{SLIDE_H}" cy="{SLIDE_W}"/></p:presentation>')
    presentation_rels = _rels_xml(
        [("rId1", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster",
          "slideMasters/slideMaster1.xml")]
        + [(f"rId{1 + i}", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide",
            f"slides/slide{i}.xml") for i in range(1, n + 1)]
        + [(f"rId{n + 2}", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme",
            "theme/theme1.xml")])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", ROOT_RELS)
        z.writestr("docProps/core.xml", CORE_PROPS)
        z.writestr("docProps/app.xml", APP_PROPS)
        z.writestr("ppt/presentation.xml", presentation)
        z.writestr("ppt/_rels/presentation.xml.rels", presentation_rels)
        z.writestr("ppt/theme/theme1.xml", THEME)
        z.writestr("ppt/slideMasters/slideMaster1.xml", SLIDE_MASTER)
        z.writestr("ppt/slideMasters/_rels/slideMaster1.xml.rels", SLIDE_MASTER_RELS)
        z.writestr("ppt/slideLayouts/slideLayout1.xml", SLIDE_LAYOUT)
        z.writestr("ppt/slideLayouts/_rels/slideLayout1.xml.rels", SLIDE_LAYOUT_RELS)

        for i, slide in enumerate(slides, start=1):
            rels: List[Tuple[str, str, str]] = [
                ("rId1", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout",
                 "../slideLayouts/slideLayout1.xml")]
            pic_xml, body_height = "", 5029200
            image_rel = slide.get("image")
            if image_rel:
                image_path = (md_path.parent / image_rel).resolve()
                if image_path.exists():
                    ext = image_path.suffix.lower().lstrip(".") or "png"
                    media_name = f"image{i}.{ext}"
                    rels.append(("rId2",
                                 "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                                 f"../media/{media_name}"))
                    width_px, height_px = _png_size(image_path)
                    scale = min(8382000 / width_px, 2286000 / height_px)
                    cx, cy = int(width_px * scale), int(height_px * scale)
                    pic_xml = _pic_shape(4, "rId2", media_name, (SLIDE_W - cx) // 2, 4114800, cx, cy)
                    body_height = 2743200
                    z.writestr(f"ppt/media/{media_name}", image_path.read_bytes())

            title = _text_shape(2, "Title", 457200, 274320, SLIDE_W - 914400, 823000,
                                [slide["title"]], size=2400, bold=True)
            body = _text_shape(3, "Body", 457200, 1219200, SLIDE_W - 914400, body_height,
                               slide["bullets"], size=1400, bullet=True)
            z.writestr(f"ppt/slides/slide{i}.xml",
                       '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                       f'<p:sld {_NS}><p:cSld><p:spTree>{_NV_GRP}{_SPTREE_EMPTY}{title}{body}'
                       f'{pic_xml}</p:spTree></p:cSld>'
                       '<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>')
            z.writestr(f"ppt/slides/_rels/slide{i}.xml.rels", _rels_xml(rels))
    return n


def run() -> Dict[str, Any]:
    """Xuất `docs/BAO-CAO.md` → .docx, `docs/slide.md` và `docs/slide-bao-ve.md` → .pptx."""
    ensure_utf8_stdio()
    out: Dict[str, Any] = {}
    report_md, slides_md = DOCS_DIR / "BAO-CAO.md", DOCS_DIR / "slide.md"
    if report_md.exists():
        counts = build_docx(report_md, DOCS_DIR / "BAO-CAO.docx")
        out["docx"] = {"path": str(DOCS_DIR / "BAO-CAO.docx"), **counts}
        print(f"Word: docs/BAO-CAO.docx — {counts['heading']} tiêu đề, {counts['table']} bảng, "
              f"{counts['image']} hình" + (f", thiếu {counts['missing_image']} hình"
                                           if counts["missing_image"] else ""))
    else:
        print("Chưa có docs/BAO-CAO.md — chạy `python -m scripts.make_report` trước.")
    if slides_md.exists():
        n = build_pptx(slides_md, DOCS_DIR / "BAO-CAO-slide.pptx")
        out["pptx"] = {"path": str(DOCS_DIR / "BAO-CAO-slide.pptx"), "slides": n}
        print(f"Slide: docs/BAO-CAO-slide.pptx — {n} slide")
    # Dàn slide BẢO VỆ (viết tay, 11 slide, có lời thoại trong file .md) — nếu có thì xuất thêm.
    defense_md = DOCS_DIR / "slide-bao-ve.md"
    if defense_md.exists():
        try:
            n_defense = build_pptx(defense_md, DOCS_DIR / "BAO-CAO-slide-bao-ve.pptx")
            out["pptx_defense"] = {"path": str(DOCS_DIR / "BAO-CAO-slide-bao-ve.pptx"),
                                   "slides": n_defense}
            print(f"Slide bảo vệ: docs/BAO-CAO-slide-bao-ve.pptx — {n_defense} slide")
        except Exception as exc:  # noqa: BLE001 - không để bước phụ làm hỏng cả bước xuất
            print(f"Bỏ qua slide bảo vệ: {exc}")
    # Bản WORD của (a) dàn slide bảo vệ và (b) bộ tài liệu bảo vệ: để đọc/duyệt và DỰNG SLIDE trực tiếp
    # (Word giữ được tiêu đề/bảng/bullet/lời thoại; mở được bằng mọi máy, không cần PowerPoint).
    for md_name, docx_name, what in (("slide-bao-ve.md", "BAO-CAO-slide-bao-ve.docx",
                                       "dàn slide bảo vệ"),
                                      ("bo-tai-lieu-bao-ve.md", "bo-tai-lieu-bao-ve.docx",
                                       "bộ tài liệu bảo vệ (factsheet + dàn slide + Q&A)")):
        src_md = DOCS_DIR / md_name
        if not src_md.exists():
            continue
        try:
            counts = build_docx(src_md, DOCS_DIR / docx_name)
            out[docx_name] = {"path": str(DOCS_DIR / docx_name), **counts}
            print(f"Word ({what}): docs/{docx_name} — {counts['heading']} tiêu đề, "
                  f"{counts['table']} bảng, {counts['image']} hình"
                  + (f", thiếu {counts['missing_image']} hình" if counts["missing_image"] else ""))
        except Exception as exc:  # noqa: BLE001 - bước phụ
            print(f"Bỏ qua {docx_name}: {exc}")
    return out


def main(argv=None) -> int:
    ensure_utf8_stdio()
    _ = argv
    print("=== Xuất Word / Slide ===")
    run()
    return 0


if __name__ == "__main__":
    sys.exit(main())